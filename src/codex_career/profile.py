from __future__ import annotations

import hashlib
import json
import re
import shutil
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader

from .db import Database, now
from .ids import new_id


def extract_resume_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        document = Document(path)
        parts = [p.text for p in document.paragraphs]
        parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        return "\n".join(part.strip() for part in parts if part.strip())
    if suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages).strip()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8")
    raise ValueError("Resume must be DOCX, PDF, TXT, or Markdown")


def evidence_id(kind: str, statement: str) -> str:
    value = f"{kind}:{' '.join(statement.lower().split())}".encode()
    return hashlib.sha256(value).hexdigest()[:16]


def normalize_profile(payload: dict[str, Any]) -> dict[str, Any]:
    required = {"name", "email", "target_roles", "locations", "work_authorization"}
    missing = sorted(required - payload.keys())
    if missing:
        raise ValueError(f"Missing profile fields: {', '.join(missing)}")
    result = dict(payload)
    result.setdefault("phone", "")
    result.setdefault("links", [])
    result.setdefault("skills", [])
    result.setdefault("preferences", {})
    result.setdefault("experience", [])
    result.setdefault("education", [])
    result.setdefault("certifications", [])
    result.setdefault("licenses", [])
    result.setdefault("projects", [])
    result.setdefault("awards", [])
    result.setdefault("accomplishments", [])
    return result


def evidence_records(profile: dict[str, Any]) -> list[tuple[str, str, list[str]]]:
    records: list[tuple[str, str, list[str]]] = []
    for kind in ("skills", "certifications", "licenses", "awards", "accomplishments"):
        for item in profile.get(kind, []):
            statement = item if isinstance(item, str) else str(item.get("name") or item.get("statement") or "")
            tags = [] if isinstance(item, str) else [str(tag) for tag in item.get("tags", [])]
            if statement:
                records.append((kind, statement, tags))
    for kind in ("experience", "projects", "leadership"):
        for item in profile.get(kind, []):
            tags = [str(tag) for tag in item.get("tags", [])]
            for bullet in item.get("bullets", item.get("accomplishments", [])):
                if bullet:
                    records.append((kind, str(bullet), tags))
    for item in profile.get("education", []):
        if isinstance(item, str):
            statement = item
            tags: list[str] = []
        else:
            statement = " | ".join(str(item.get(key, "")).strip() for key in ("degree", "school", "dates") if item.get(key))
            tags = [str(tag) for tag in item.get("tags", [])]
        if statement:
            records.append(("education", statement, tags))
    return records


def save_profile(db: Database, payload: dict[str, Any], source: str = "onboarding") -> dict[str, Any]:
    profile = normalize_profile(payload)
    timestamp = now()
    with db.connect() as connection:
        version = connection.execute("SELECT COALESCE(MAX(version), 0) + 1 FROM profile_versions").fetchone()[0]
        connection.execute(
            "INSERT INTO profile_versions(public_id, version, data_json, source_resume_name, reviewed_at, created_at) VALUES(?,?,?,?,?,?)",
            (new_id("prf"), version, json.dumps(profile, sort_keys=True), source, timestamp, timestamp),
        )
        connection.execute(
            "INSERT INTO profile(id, data_json, updated_at) VALUES(1,?,?) "
            "ON CONFLICT(id) DO UPDATE SET data_json=excluded.data_json, updated_at=excluded.updated_at",
            (json.dumps(profile, sort_keys=True), timestamp),
        )
        connection.execute("DELETE FROM evidence")
        vault = {
            "name": profile.get("name"), "email": profile.get("email"), "phone": profile.get("phone"),
            "location": profile.get("locations"), "work_authorization": profile.get("work_authorization"),
            "education": profile.get("education"), "certifications": profile.get("certifications"),
        }
        for key, value in vault.items():
            if value in (None, "", []):
                continue
            connection.execute(
                "INSERT INTO answer_vault(public_id, answer_key, value_json, classification, verified, provenance, updated_at) VALUES(?,?,?,?,1,?,?) "
                "ON CONFLICT(answer_key) DO UPDATE SET value_json=excluded.value_json, verified=1, provenance=excluded.provenance, updated_at=excluded.updated_at",
                (new_id("ans"), key, json.dumps(value, sort_keys=True), "AUTO", source, timestamp),
            )
        for kind, statement, tags in evidence_records(profile):
            connection.execute(
                "INSERT OR IGNORE INTO evidence(id, public_id, kind, source, statement, verified, provenance_json, tags_json, created_at, updated_at) VALUES(?,?,?,?,?,1,?,?,?,?)",
                (evidence_id(kind, statement), new_id("evd"), kind, source, statement, json.dumps({"profile_version": version, "source": source}, sort_keys=True), json.dumps(tags, sort_keys=True), timestamp, timestamp),
            )
    return profile


def preserve_resume(source: Path, destination_root: Path) -> Path:
    destination = destination_root / "resume" / f"original{source.suffix.lower()}"
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    return destination


def load_profile(db: Database) -> dict[str, Any]:
    with db.connect() as connection:
        row = connection.execute("SELECT data_json FROM profile WHERE id=1").fetchone()
    if not row:
        raise ValueError("No profile found. Run `job-agent setup` first.")
    return json.loads(row["data_json"])


def resume_hints(text: str) -> dict[str, Any]:
    email = re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", text)
    phone = re.search(r"(?:\+?1[ .-]?)?\(?\d{3}\)?[ .-]\d{3}[ .-]\d{4}", text)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return {
        "name": lines[0] if lines else "",
        "email": email.group(0) if email else "",
        "phone": phone.group(0) if phone else "",
        "resume_text": text,
    }


def onboarding_packet(path: Path) -> dict[str, Any]:
    text = extract_resume_text(path)
    return {
        "instructions": [
            "Create a profile draft using only facts explicitly present in resume_text.",
            "Leave uncertain fields empty and put follow-up needs in questions_for_user.",
            "The user must review every claim before importing this profile.",
        ],
        "resume_file": path.name,
        "resume_text": text,
        "detected": resume_hints(text),
        "response_schema": {
            "name": "string",
            "email": "string",
            "phone": "string",
            "links": ["string"],
            "target_roles": ["string from user answers"],
            "locations": ["string from user answers"],
            "work_authorization": "string from user answer",
            "skills": ["verified skill"],
            "preferences": {},
            "experience": [{"company": "string", "title": "string", "dates": "string", "bullets": ["verbatim fact"]}],
            "education": [{"degree": "string", "school": "string", "dates": "string"}],
            "certifications": ["string"],
            "licenses": ["string"],
            "projects": [{"name": "string", "dates": "string", "bullets": ["verbatim fact"], "tags": ["string"]}],
            "awards": ["string"],
            "accomplishments": ["string"],
            "questions_for_user": ["unresolved onboarding question"],
        },
    }
