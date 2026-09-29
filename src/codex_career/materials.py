from __future__ import annotations

import hashlib
import html
import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate

from .db import Database, now
from .ids import new_id

TAILORING_MODES = {"LIGHT", "STANDARD", "DEEP"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _numbers(text: str) -> set[str]:
    return set(re.findall(r"(?:[$+]\s*)?\d[\d,.]*(?:\s*%|\+)?", text))


def validate_proposal(db: Database, proposal: dict[str, Any]) -> dict[str, Any]:
    mode = str(proposal.get("tailoring_mode", "STANDARD")).upper()
    if mode not in TAILORING_MODES:
        raise ValueError("tailoring_mode must be LIGHT, STANDARD, or DEEP")
    evidence_ids = set(proposal.get("selected_evidence", []))
    rewrites = proposal.get("bullet_rewrites", [])
    evidence_ids.update(item.get("evidence_id") for item in rewrites)
    evidence_ids.discard(None)
    if not evidence_ids:
        raise ValueError("Tailoring proposal must select verified evidence")
    placeholders = ",".join("?" for _ in evidence_ids)
    with db.connect() as connection:
        rows = connection.execute(
            f"SELECT id, statement, verified FROM evidence WHERE id IN ({placeholders})", tuple(evidence_ids)
        ).fetchall()
    found = {row["id"]: row for row in rows}
    missing = sorted(evidence_ids - found.keys())
    if missing:
        raise ValueError(f"Unknown evidence IDs: {', '.join(missing)}")
    unverified = sorted(key for key, row in found.items() if not row["verified"])
    if unverified:
        raise ValueError(f"Unverified evidence IDs: {', '.join(unverified)}")
    for item in rewrites:
        original = found[item["evidence_id"]]["statement"]
        invented_numbers = _numbers(item["text"]) - _numbers(original)
        if invented_numbers:
            raise ValueError(
                f"Rewrite for {item['evidence_id']} introduced numeric claims: {', '.join(sorted(invented_numbers))}"
            )
    allowed_numbers = set().union(*(_numbers(row["statement"]) for row in found.values()))
    for field in ("summary", "cover_letter"):
        invented_numbers = _numbers(proposal.get(field, "")) - allowed_numbers
        if invented_numbers:
            raise ValueError(f"{field} introduced numeric claims: {', '.join(sorted(invented_numbers))}")
    return proposal


def tailoring_packet(db: Database, job_id: int, profile: dict[str, Any], mode: str = "STANDARD") -> dict[str, Any]:
    mode = mode.upper()
    if mode not in TAILORING_MODES:
        raise ValueError("Tailoring mode must be LIGHT, STANDARD, or DEEP")
    with db.connect() as connection:
        job = connection.execute("SELECT * FROM jobs WHERE id=?", (job_id,)).fetchone()
        evidence = connection.execute(
            "SELECT id, kind, statement FROM evidence WHERE verified=1 ORDER BY kind, id"
        ).fetchall()
    if not job:
        raise ValueError(f"Job {job_id} not found")
    return {
        "instructions": [
            "Use only the verified evidence supplied below.",
            "Do not add employers, dates, credentials, technologies, metrics, or responsibilities.",
            "Return JSON matching the response_schema. Keep rewrites factual and ATS-readable.",
        ],
        "tailoring_mode": mode,
        "job": {key: job[key] for key in ("id", "company", "title", "location", "description", "url")},
        "candidate": {
            "target_roles": profile.get("target_roles", []),
            "skills": profile.get("skills", []),
            "evidence": [dict(row) for row in evidence],
        },
        "response_schema": {
            "tailoring_mode": mode,
            "summary": "string",
            "selected_evidence": ["evidence id"],
            "bullet_rewrites": [{"evidence_id": "id", "text": "truthful rewrite"}],
            "cover_letter": "string",
        },
    }


def _configure_styles(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(2)
    for name, size in (("Title", 18), ("Heading 1", 11.5), ("Heading 2", 10.5)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(5)
        style.paragraph_format.space_after = Pt(2)
        border = style.element.pPr.find(qn("w:pBdr"))
        if border is not None:
            style.element.pPr.remove(border)
    if "Resume Bullet" not in [style.name for style in document.styles]:
        bullet = document.styles.add_style("Resume Bullet", WD_STYLE_TYPE.PARAGRAPH)
        bullet.base_style = document.styles["Normal"]
        bullet.paragraph_format.left_indent = Inches(0.18)
        bullet.paragraph_format.first_line_indent = Inches(-0.18)
        bullet.paragraph_format.space_after = Pt(1)


def _selected_content(profile: dict[str, Any], proposal: dict[str, Any]) -> tuple[list[str], list[tuple[str, dict[str, Any], list[tuple[str, str]]]]]:
    from .profile import evidence_id

    selected = set(proposal.get("selected_evidence", []))
    rewrites = {item["evidence_id"]: item["text"] for item in proposal.get("bullet_rewrites", [])}
    selected_skills = [skill for skill in profile.get("skills", []) if evidence_id("skills", str(skill)) in selected]
    skills = selected_skills or [str(skill) for skill in profile.get("skills", [])]
    sections = []
    for kind, heading in (("experience", "Experience"), ("projects", "Projects"), ("leadership", "Leadership")):
        for item in profile.get(kind, []):
            bullets = []
            for bullet in item.get("bullets", []):
                item_id = evidence_id(kind, bullet)
                if item_id in selected:
                    bullets.append((item_id, rewrites.get(item_id, bullet)))
            if bullets:
                sections.append((heading, item, bullets))
    return skills, sections


def render_resume(path: Path, profile: dict[str, Any], proposal: dict[str, Any]) -> None:
    document = Document()
    _configure_styles(document)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(profile["name"])
    contact_values = [profile.get("email", ""), profile.get("phone", ""), *profile.get("links", [])]
    contact = document.add_paragraph(" | ".join(value for value in contact_values if value))
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if proposal.get("summary"):
        document.add_heading("Professional Summary", level=1)
        document.add_paragraph(proposal["summary"])
    skills, sections = _selected_content(profile, proposal)
    if skills:
        document.add_heading("Skills", level=1)
        document.add_paragraph(" | ".join(skills))
    current_heading = None
    for section_heading, role, bullets in sections:
        if section_heading != current_heading:
            document.add_heading(section_heading, level=1)
            current_heading = section_heading
        heading = document.add_paragraph(style="Heading 2")
        heading.add_run(role.get("title") or role.get("name") or role.get("organization", "")).bold = True
        company = role.get("company") or role.get("organization", "")
        dates = role.get("dates", "")
        if company and company not in heading.text:
            heading.add_run(f" | {company}")
        if dates:
            heading.add_run(f" | {dates}")
        for _, text in bullets:
            document.add_paragraph(f"• {text}", style="Resume Bullet")
    if profile.get("education"):
        document.add_heading("Education", level=1)
        for item in profile["education"]:
            if isinstance(item, str):
                document.add_paragraph(item)
            else:
                document.add_paragraph(" | ".join(str(item.get(key, "")) for key in ("degree", "school", "dates") if item.get(key)))
    if profile.get("certifications"):
        document.add_heading("Certifications", level=1)
        document.add_paragraph(" | ".join(str(item) for item in profile["certifications"]))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def render_resume_pdf(path: Path, profile: dict[str, Any], proposal: dict[str, Any]) -> None:
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "Resume Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=10.5, leading=13, spaceAfter=2
    )
    title = ParagraphStyle(
        "Resume Title", parent=body, fontName="Helvetica-Bold", fontSize=18, leading=20, alignment=TA_CENTER, spaceAfter=3
    )
    contact = ParagraphStyle("Resume Contact", parent=body, alignment=TA_CENTER, spaceAfter=6)
    heading = ParagraphStyle(
        "Resume Heading", parent=body, fontName="Helvetica-Bold", fontSize=11.5, leading=14, spaceBefore=5, spaceAfter=2
    )
    role = ParagraphStyle(
        "Resume Role", parent=body, fontName="Helvetica-Bold", fontSize=10.5, leading=13, spaceBefore=2, spaceAfter=1
    )
    bullet = ParagraphStyle("Resume Bullet", parent=body, leftIndent=12, firstLineIndent=-8, bulletIndent=0, spaceAfter=1)
    story = [Paragraph(html.escape(profile["name"]), title)]
    contact_values = [profile.get("email", ""), profile.get("phone", ""), *profile.get("links", [])]
    story.append(Paragraph(html.escape(" | ".join(value for value in contact_values if value)), contact))
    if proposal.get("summary"):
        story.extend([Paragraph("Professional Summary", heading), Paragraph(html.escape(proposal["summary"]), body)])
    skills, sections = _selected_content(profile, proposal)
    if skills:
        story.extend(
            [Paragraph("Skills", heading), Paragraph(html.escape(" | ".join(skills)), body)]
        )
    current_heading = None
    for section_heading, item, bullets in sections:
        if section_heading != current_heading:
            story.append(Paragraph(section_heading, heading))
            current_heading = section_heading
        label = " | ".join(str(value) for value in (item.get("title") or item.get("name"), item.get("company") or item.get("organization"), item.get("dates")) if value)
        story.append(Paragraph(html.escape(label), role))
        for _, text in bullets:
            story.append(Paragraph(f"• {html.escape(text)}", bullet))
    if profile.get("education"):
        story.append(Paragraph("Education", heading))
        for item in profile["education"]:
            text = item if isinstance(item, str) else " | ".join(
                str(item.get(key, "")) for key in ("degree", "school", "dates") if item.get(key)
            )
            story.append(Paragraph(html.escape(text), body))
    if profile.get("certifications"):
        story.extend(
            [
                Paragraph("Certifications", heading),
                Paragraph(html.escape(" | ".join(str(item) for item in profile["certifications"])), body),
            ]
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(
        str(path), pagesize=letter, leftMargin=0.65 * inch, rightMargin=0.65 * inch, topMargin=0.55 * inch, bottomMargin=0.55 * inch
    ).build(story)


TEMPLATES = {"classic-ats": (1, render_resume, render_resume_pdf)}


def validate_documents(
    docx_path: Path, pdf_path: Path | None = None, *, profile: dict[str, Any] | None = None, target_pages: int = 1
) -> dict[str, Any]:
    document = Document(docx_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if not text.strip():
        raise ValueError("Generated DOCX contains no text")
    expected = ["Experience"]
    if profile:
        if profile.get("skills"):
            expected.append("Skills")
        if profile.get("education"):
            expected.append("Education")
        if profile.get("certifications"):
            expected.append("Certifications")
        for value in (profile.get("name"), profile.get("email")):
            if value and str(value) not in text:
                raise ValueError(f"Generated DOCX is missing required contact value: {value}")
    missing_sections = [section for section in expected if section not in text]
    if missing_sections:
        raise ValueError(f"Generated DOCX is missing sections: {', '.join(missing_sections)}")
    broken_bullets = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name == "Resume Bullet" and not paragraph.text.removeprefix("•").strip()]
    if broken_bullets:
        raise ValueError("Generated DOCX contains an empty resume bullet")
    def style_size(style: Any) -> float | None:
        while style is not None:
            if getattr(style, "font", None) and style.font.size:
                return style.font.size.pt
            style = getattr(style, "base_style", None)
        return None

    font_sizes = [size for paragraph in document.paragraphs if (size := style_size(paragraph.style)) is not None]
    minimum_font = min(font_sizes) if font_sizes else None
    if minimum_font is not None and minimum_font < 9:
        raise ValueError(f"Generated DOCX uses a font below 9pt: {minimum_font}")
    result: dict[str, Any] = {
        "docx_sha256": sha256(docx_path), "docx_paragraphs": len(document.paragraphs),
        "required_sections": expected, "minimum_style_font_pt": minimum_font, "empty_bullets": 0,
    }
    if pdf_path:
        reader = PdfReader(pdf_path)
        page_text = [(page.extract_text() or "").strip() for page in reader.pages]
        pdf_text = "\n".join(page_text)
        if not pdf_text.strip():
            raise ValueError("Generated PDF contains no extractable text")
        if any(not item for item in page_text):
            raise ValueError("Generated PDF contains a blank page")
        if len(reader.pages) > target_pages:
            raise ValueError(
                f"Generated PDF is {len(reader.pages)} pages; target is {target_pages}. Shorten weak content before changing typography."
            )
        missing_pdf_sections = [section for section in expected if section not in pdf_text]
        if missing_pdf_sections:
            raise ValueError(f"Generated PDF is missing sections: {', '.join(missing_pdf_sections)}")
        result.update(pdf_sha256=sha256(pdf_path), pdf_pages=len(reader.pages), blank_pdf_pages=0, target_pages=target_pages)
    return result


def build_materials(
    db: Database,
    root: Path,
    job_id: int,
    profile: dict[str, Any],
    proposal: dict[str, Any],
    *,
    make_pdf: bool = False,
) -> dict[str, Any]:
    validate_proposal(db, proposal)
    template_id = str(proposal.get("template_id", "classic-ats"))
    try:
        template_version, docx_renderer, pdf_renderer = TEMPLATES[template_id]
    except KeyError as error:
        raise ValueError(f"Unknown resume template: {template_id}") from error
    with db.connect() as connection:
        job = connection.execute("SELECT * FROM jobs WHERE id=?", (job_id,)).fetchone()
        if not job:
            raise ValueError(f"Job {job_id} not found")
        version = connection.execute(
            "SELECT COALESCE(MAX(version), 0) + 1 AS version FROM material_sets WHERE job_id=?", (job_id,)
        ).fetchone()["version"]
    folder = root / f"job-{job_id}" / f"v{version}"
    folder.mkdir(parents=True, exist_ok=False)
    resume_path = folder / "resume.docx"
    docx_renderer(resume_path, profile, proposal)
    (folder / "profile.json").write_text(json.dumps(profile, indent=2, sort_keys=True), encoding="utf-8")
    (folder / "job.json").write_text(json.dumps(dict(job), indent=2, sort_keys=True), encoding="utf-8")
    (folder / "proposal.json").write_text(json.dumps(proposal, indent=2, sort_keys=True), encoding="utf-8")
    if proposal.get("cover_letter"):
        (folder / "cover-letter.txt").write_text(proposal["cover_letter"].strip() + "\n", encoding="utf-8")
    pdf_path = folder / "resume.pdf" if make_pdf else None
    if pdf_path:
        pdf_renderer(pdf_path, profile, proposal)
    validation = validate_documents(
        resume_path, pdf_path, profile=profile, target_pages=int(proposal.get("target_resume_pages", 1))
    )
    files = []
    for path in sorted(folder.iterdir()):
        if path.is_file():
            files.append({"name": path.name, "sha256": sha256(path), "bytes": path.stat().st_size})
    manifest = {
        "job_id": job_id,
        "version": version,
        "template_id": template_id,
        "template_version": template_version,
        "tailoring_mode": str(proposal.get("tailoring_mode", "STANDARD")).upper(),
        "created_at": now(),
        "files": files,
        "validation": validation,
    }
    manifest_path = folder / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, sort_keys=True), encoding="utf-8")
    with db.connect() as connection:
        profile_version = connection.execute("SELECT id FROM profile_versions ORDER BY version DESC LIMIT 1").fetchone()
        material_public_id = new_id("mat")
        cursor = connection.execute(
            "INSERT INTO material_sets(public_id, job_id, version, status, manifest_json, profile_version_id, template_id, template_version, tailoring_mode, structured_json, created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
            (material_public_id, job_id, version, "draft", json.dumps({**manifest, "path": str(folder)}, sort_keys=True), profile_version["id"] if profile_version else None, template_id, template_version, manifest["tailoring_mode"].lower(), json.dumps(proposal, sort_keys=True), now()),
        )
        material_id = cursor.lastrowid
        for item in files:
            kind = "resume" if item["name"].startswith("resume.") else "cover_letter" if item["name"].startswith("cover-letter") else "source_snapshot"
            document_id = connection.execute(
                "INSERT INTO documents(public_id, kind, created_at) VALUES(?,?,?) RETURNING id",
                (new_id("doc"), kind, now()),
            ).fetchone()["id"]
            connection.execute(
                "INSERT INTO document_versions(public_id, document_id, version, job_id, profile_version_id, template_id, template_version, tailoring_mode, path, sha256, metadata_json, created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                (new_id("dv"), document_id, 1, job_id, profile_version["id"] if profile_version else None, template_id, template_version, manifest["tailoring_mode"], str(folder / item["name"]), item["sha256"], json.dumps({"material_set_id": material_id, "bytes": item["bytes"]}), now()),
            )
    return {"id": material_id, "public_id": material_public_id, "path": str(folder), **manifest}


def approve_materials(db: Database, material_id: int) -> None:
    verify_materials(db, material_id)
    with db.connect() as connection:
        changed = connection.execute(
            "UPDATE material_sets SET status='approved' WHERE id=? AND status='draft'", (material_id,)
        ).rowcount
    if not changed:
        raise ValueError("Material set is missing or is not a draft")
    db.event("material_set", material_id, "approved", {})


def verify_materials(db: Database, material_id: int) -> dict[str, Any]:
    with db.connect() as connection:
        row = connection.execute("SELECT manifest_json FROM material_sets WHERE id=?", (material_id,)).fetchone()
    if not row:
        raise ValueError(f"Material set {material_id} not found")
    manifest = json.loads(row["manifest_json"])
    folder = Path(manifest["path"])
    for entry in manifest["files"]:
        path = folder / entry["name"]
        if not path.is_file() or sha256(path) != entry["sha256"]:
            raise ValueError(f"Material integrity check failed: {entry['name']}")
    return manifest
